# -*- coding: utf-8 -*-
"""Builds the Qwader manual test-case workbook (Arabic, right-to-left)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit('\\', 1)[0] if '\\' in __file__ else '.')
from cases import CASES, MODULES  # noqa: E402

OUT = sys.argv[1]
FONT = 'Arial'
BRAND = '0F5F5A'
INPUT_FILL = PatternFill('solid', fgColor='FFF7D6')   # cells the tester fills in
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
BAND_FILL = PatternFill('solid', fgColor='F2F6F6')
thin = Side(style='thin', color='D0D7D7')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
CENTER = Alignment(wrap_text=True, vertical='top', horizontal='center', readingOrder=2)

STATUSES = ['لم يُختبر', 'نجح', 'فشل', 'محجوب']
module_name = dict(MODULES)
order = {code: i for i, (code, _) in enumerate(MODULES)}
cases = sorted(CASES, key=lambda c: order[c[0]])  # stable: keeps the written order inside a module
assert all(c[0] in module_name for c in cases), 'unknown module'

wb = Workbook()

# ------------------------------------------------------------------ test cases
ws = wb.active
ws.title = 'حالات الاختبار'
ws.sheet_view.rightToLeft = True
headers = ['رقم الحالة', 'القسم', 'السيناريو', 'الدور', 'المتطلبات المسبقة', 'خطوات التنفيذ', 'النتيجة المتوقعة',
           'الأولوية', 'مغطى آليًا', 'الاختبار الآلي', 'النتيجة', 'النتيجة الفعلية / الملاحظات', 'المُختبِر', 'تاريخ الاختبار']
widths = [11, 22, 34, 14, 26, 46, 46, 9, 9, 30, 12, 34, 14, 14]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    cell = ws.cell(row=1, column=i, value=h)
    cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30

counters = {}
for r, (code, scenario, role, pre, steps, expected, priority, auto) in enumerate(cases, start=2):
    counters[code] = counters.get(code, 0) + 1
    values = [f'{code}-{counters[code]:02d}', module_name[code], scenario, role, pre, steps, expected, priority,
              'نعم' if auto else 'لا', auto or '—', 'لم يُختبر', None, None, None]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, size=10, color='6B7B7B' if c == 10 else '1F2A2A')
        cell.alignment = CENTER if c in (1, 4, 8, 9, 11, 13, 14) else WRAP
        cell.border = BORDER
        if c >= 11:
            cell.fill = INPUT_FILL
    if c == 14:
        ws.cell(row=r, column=14).number_format = 'yyyy-mm-dd'
last = len(cases) + 1

ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{last}'

dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True,
                    showErrorMessage=True, errorTitle='قيمة غير صحيحة', error='اختر من القائمة: ' + '، '.join(STATUSES))
ws.add_data_validation(dv)
dv.add(f'K2:K{last}')

status_colors = {'نجح': ('DCF5E7', '146C43'), 'فشل': ('FDE2E1', 'B42318'), 'محجوب': ('FFF1D6', '9A6700')}
for status, (bg, fg) in status_colors.items():
    ws.conditional_formatting.add(f'K2:K{last}', FormulaRule(formula=[f'$K2="{status}"'],
                                  fill=PatternFill('solid', fgColor=bg), font=Font(name=FONT, bold=True, color=fg)))
ws.conditional_formatting.add(f'H2:H{last}', FormulaRule(formula=['$H2="عالية"'], font=Font(name=FONT, bold=True, color='B42318')))

# ------------------------------------------------------------------ summary
ss = wb.create_sheet('الملخص')
ss.sheet_view.rightToLeft = True
ss['A1'] = 'ملخص نتائج الاختبار — نظام معرض قويدر'
ss['A1'].font = Font(name=FONT, bold=True, size=14, color=BRAND)
ss['A2'] = 'يتحدث تلقائيًا من عمود «النتيجة» في ورقة «حالات الاختبار».'
ss['A2'].font = Font(name=FONT, size=10, color='6B7B7B')
sum_headers = ['القسم', 'عدد الحالات', 'نجح', 'فشل', 'محجوب', 'لم يُختبر', 'مغطى آليًا', 'نسبة الإنجاز', 'نسبة النجاح']
for i, h in enumerate(sum_headers, start=1):
    cell = ss.cell(row=4, column=i, value=h)
    cell.font = Font(name=FONT, bold=True, color='FFFFFF')
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER
for i, w in enumerate([30, 12, 10, 10, 10, 12, 12, 13, 13], start=1):
    ss.column_dimensions[get_column_letter(i)].width = w

TC = "'حالات الاختبار'"
rng = lambda col: f'{TC}!${col}$2:${col}${last}'  # noqa: E731
row = 5
for code, name in MODULES:
    ss.cell(row=row, column=1, value=name)
    ss.cell(row=row, column=2, value=f'=COUNTIF({rng("B")},A{row})')
    for col, status in zip((3, 4, 5), ('نجح', 'فشل', 'محجوب')):
        ss.cell(row=row, column=col, value=f'=COUNTIFS({rng("B")},$A{row},{rng("K")},"{status}")')
    ss.cell(row=row, column=6, value=f'=B{row}-C{row}-D{row}-E{row}')
    ss.cell(row=row, column=7, value=f'=COUNTIFS({rng("B")},$A{row},{rng("I")},"نعم")')
    ss.cell(row=row, column=8, value=f'=IF(B{row}=0,0,(C{row}+D{row}+E{row})/B{row})')
    ss.cell(row=row, column=9, value=f'=IF(C{row}+D{row}=0,0,C{row}/(C{row}+D{row}))')
    row += 1
first, end = 5, row - 1
ss.cell(row=row, column=1, value='الإجمالي')
for col in range(2, 8):
    L_ = get_column_letter(col)
    ss.cell(row=row, column=col, value=f'=SUM({L_}{first}:{L_}{end})')
ss.cell(row=row, column=8, value=f'=IF(B{row}=0,0,(C{row}+D{row}+E{row})/B{row})')
ss.cell(row=row, column=9, value=f'=IF(C{row}+D{row}=0,0,C{row}/(C{row}+D{row}))')
total_row = row
for r in range(first, total_row + 1):
    for c in range(1, 10):
        cell = ss.cell(row=r, column=c)
        cell.border = BORDER
        cell.alignment = WRAP if c == 1 else CENTER
        cell.font = Font(name=FONT, size=10, bold=(r == total_row))
        if c in (8, 9):
            cell.number_format = '0%'
        if r == total_row:
            cell.fill = BAND_FILL

# by priority
row = total_row + 3
ss.cell(row=row, column=1, value='حسب الأولوية').font = Font(name=FONT, bold=True, size=12, color=BRAND)
row += 1
for i, h in enumerate(['الأولوية', 'عدد الحالات', 'نجح', 'فشل', 'لم يُختبر'], start=1):
    cell = ss.cell(row=row, column=i, value=h)
    cell.font = Font(name=FONT, bold=True, color='FFFFFF')
    cell.fill = HEAD_FILL
    cell.alignment = CENTER
    cell.border = BORDER
for p in ('عالية', 'متوسطة', 'منخفضة'):
    row += 1
    ss.cell(row=row, column=1, value=p)
    ss.cell(row=row, column=2, value=f'=COUNTIF({rng("H")},A{row})')
    ss.cell(row=row, column=3, value=f'=COUNTIFS({rng("H")},$A{row},{rng("K")},"نجح")')
    ss.cell(row=row, column=4, value=f'=COUNTIFS({rng("H")},$A{row},{rng("K")},"فشل")')
    ss.cell(row=row, column=5, value=f'=COUNTIFS({rng("H")},$A{row},{rng("K")},"لم يُختبر")')
    for c in range(1, 6):
        cell = ss.cell(row=row, column=c)
        cell.border = BORDER
        cell.alignment = CENTER
        cell.font = Font(name=FONT, size=10)
row += 2
ss.cell(row=row, column=1, value='جاهز للتشغيل الفعلي عندما: كل حالات الأولوية «عالية» = نجح، ولا يوجد «فشل» مفتوح.').font = Font(name=FONT, size=10, bold=True)

# ------------------------------------------------------------------ guide
gd = wb.create_sheet('دليل الاستخدام', 0)
gd.sheet_view.rightToLeft = True
gd.column_dimensions['A'].width = 26
gd.column_dimensions['B'].width = 90
lines = [
    ('حالات اختبار نظام معرض قويدر', None, 'title'),
    ('الهدف', 'التأكد من أن كل أقسام النظام تعمل كما هو متوقع قبل الاعتماد عليه في العمل اليومي.', None),
    ('عدد الحالات', f'{len(cases)} حالة في {len(MODULES)} قسمًا.', None),
    ('طريقة الاستخدام', '1) افتح ورقة «حالات الاختبار».\n2) نفّذ خطوات كل حالة بالدور المذكور.\n3) اختر النتيجة من القائمة واكتب ما حدث فعلًا إذا اختلف.\n4) راجع ورقة «الملخص» لمعرفة نسبة الإنجاز والنجاح.', None),
    ('الخانات التي تعبّئها', 'الخانات الصفراء فقط: النتيجة، النتيجة الفعلية / الملاحظات، المُختبِر، تاريخ الاختبار. باقي الأعمدة ثابتة.', 'input'),
    ('قيم «النتيجة»', 'نجح: النتيجة مطابقة للمتوقع.  فشل: النتيجة مختلفة (اكتب ما حدث).  محجوب: لا يمكن التنفيذ الآن (مثلًا لا يوجد آيفون).  لم يُختبر: لم تُنفّذ بعد.', None),
    ('مثال تعبئة', 'النتيجة: فشل  |  النتيجة الفعلية: «الفاتورة بقيت آجل بعد تأكيد الحوالة»  |  المُختبِر: محمد  |  التاريخ: 2026-10-26', 'input'),
    ('مغطى آليًا', '«نعم» تعني أن هذه الحالة يفحصها أيضًا اختبار آلي يعمل مع كل تعديل على الكود (اسمه في عمود «الاختبار الآلي»). الاختبار اليدوي يبقى مطلوبًا للتأكد من الشاشات على أجهزتك.', None),
    ('الأدوار', 'مدير النظام، مدير المعرض، موظف مبيعات، أمين مخزن، محاسب، مشاهد. أنشئ مستخدمًا تجريبيًا لكل دور قبل البدء.', None),
    ('تحذير', 'حالات «تصفير النظام» تمسح البيانات نهائيًا: نفّذها فقط على نسخة تجريبية أو قبل إدخال بيانات حقيقية.', 'warn'),
    ('ترقيم الحالات', 'رمز القسم + رقم، مثل POS-03 = الحالة الثالثة في نقطة البيع.', None),
]
r = 1
for label, text, kind in lines:
    a = gd.cell(row=r, column=1, value=label)
    if kind == 'title':
        a.font = Font(name=FONT, bold=True, size=16, color=BRAND)
        gd.row_dimensions[r].height = 28
        r += 2
        continue
    a.font = Font(name=FONT, bold=True, size=11, color=BRAND)
    a.alignment = WRAP
    b = gd.cell(row=r, column=2, value=text)
    b.font = Font(name=FONT, size=11, color='B42318' if kind == 'warn' else '1F2A2A', bold=(kind == 'warn'))
    b.alignment = WRAP
    if kind == 'input':
        b.fill = INPUT_FILL
    gd.row_dimensions[r].height = max(22, 18 * (text.count('\n') + 1 + len(text) // 95))
    r += 1

r += 1
gd.cell(row=r, column=1, value='الأقسام').font = Font(name=FONT, bold=True, size=12, color=BRAND)
r += 1
for code, name in MODULES:
    gd.cell(row=r, column=1, value=code).font = Font(name=FONT, size=10, bold=True)
    gd.cell(row=r, column=1).alignment = CENTER
    gd.cell(row=r, column=2, value=f'{name} — {sum(1 for c in cases if c[0] == code)} حالة').font = Font(name=FONT, size=10)
    gd.cell(row=r, column=2).alignment = WRAP
    r += 1

wb.active = 1
wb.save(OUT)
print(f'saved {OUT}: {len(cases)} cases')
